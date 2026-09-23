"""Exportación de reportes: el Excel conserva las cifras como números y el
historial del estudiante se arma con sus totales."""

from openpyxl import load_workbook

from pages.student_history import build_report
from widgets.report_export import Report, Section, report_html, write_pdf, write_xlsx

HISTORY = {
    "student": {"id": 7, "full_name": "Ana María Pérez", "email": "ana@correo.com", "contact_phone": "7777-8888"},
    "enrollments": [{
        "id": 3, "diploma_name": "Marketing Digital", "schedule_name": "Sábados 8:00 - 10:15",
        "tuition_plan": "GRUPAL", "status": "PENDIENTE", "enrollment_date": "2026-08-01",
        "start_date": "2026-08-01", "next_payment_date": "2026-08-29", "is_overdue": True,
    }],
    "payments": [
        {"id": 10, "enrollment_id": 3, "diploma_name": "Marketing Digital", "kind": "MATRICULA",
         "payment_date": "2026-08-01", "due_date": "2026-08-01", "amount": 20.0, "surcharge": 0.0,
         "total": 20.0, "payment_type": "Efectivo", "status": "PAGADO", "cashier_name": "Carla Méndez",
         "observations": "Matrícula completa"},
        {"id": 11, "enrollment_id": 3, "diploma_name": "Marketing Digital", "kind": "COLEGIATURA",
         "payment_date": "2026-08-01", "due_date": "2026-08-01", "amount": 25.0, "surcharge": 3.0,
         "total": 28.0, "payment_type": "Efectivo", "status": "PAGADO", "cashier_name": None,
         "observations": None},
    ],
    "totals": {"paid": 48.0, "registration": 20.0, "tuition": 28.0, "surcharges": 3.0, "voided": 0.0,
               "tuition_installments_paid": 1, "overdue_enrollments": 1},
}


def test_history_report_has_totals_and_both_sections():
    report = build_report(HISTORY)
    assert report.title == "Historial de pagos — Ana María Pérez"
    assert ("Total pagado", 48.0) in report.summary
    assert ("Cuotas de colegiatura pagadas", 1) in report.summary
    enrollments, payments = report.sections
    assert enrollments.rows[0][-1] == "2026-08-29 (atrasado)"
    assert payments.rows[1][2] == "Colegiatura"
    assert payments.rows[1][10] == "—"  # sin cajero registrado


def test_xlsx_keeps_amounts_as_numbers(tmp_path):
    path = tmp_path / "historial.xlsx"
    write_xlsx(build_report(HISTORY), str(path))
    sheet = load_workbook(path).active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values[1][0] == "Historial de pagos — Ana María Pérez"
    summary = {row[0]: row[1] for row in values if row and row[0] in ("Total pagado", "Inscripciones atrasadas")}
    assert summary == {"Total pagado": 48.0, "Inscripciones atrasadas": 1}
    payment_row = next(row for row in values if row[0] == 11)
    assert payment_row[7] == 28.0  # total como número, no como texto "$28.00"
    money_cell = next(c for c in sheet.iter_rows() if c[0].value == 11)[7]
    assert "$" in money_cell.number_format


def test_empty_section_is_marked():
    report = Report("Vacío", sections=[Section("Pagos", ["A", "B"], [])])
    assert "Sin registros" in report_html(report)


def test_pdf_is_written(tmp_path, qapp):
    path = tmp_path / "reporte.pdf"
    write_pdf(build_report(HISTORY), str(path))
    assert path.read_bytes().startswith(b"%PDF")


def test_ticket_pdf_is_a_single_80mm_roll_with_dark_text(tmp_path, qapp):
    from PySide6.QtPdf import QPdfDocument

    from widgets import ticket_printer

    path = tmp_path / "ticket.pdf"
    assert ticket_printer.write_ticket_pdf(
        str(path), "Comprobante de pago", [("Total", "$28.00"), ("Cambio", "$2.00")], "Conserve este comprobante."
    )
    document = QPdfDocument()
    document.load(str(path))
    assert document.pageCount() == 1
    width_mm = document.pagePointSize(0).width() * 25.4 / 72
    assert abs(width_mm - 80) < 1
    # Antes usaba el color de texto del tema oscuro (#e6ebff): ilegible en papel.
    html = ticket_printer._ticket_document("T", [("A", "B")], "").toHtml().lower()
    assert "#e6ebff" not in html
