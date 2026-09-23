"""Exportar reportes a Excel (.xlsx) o PDF.

Excel usa openpyxl (las cifras quedan como números, para que se puedan sumar
o filtrar en Excel). El PDF usa QTextDocument + QPrinter, que ya vienen con
PySide6, así que no agrega otra librería al EXE.
"""

from dataclasses import dataclass, field
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Any, Optional

from PySide6.QtCore import QMarginsF, QStandardPaths
from PySide6.QtGui import QPageLayout, QPageSize, QTextDocument
from PySide6.QtPrintSupport import QPrinter
from PySide6.QtWidgets import QFileDialog, QMessageBox, QWidget

INSTITUTION = "CTC El Salvador"
_MONEY_FORMAT = '"$"#,##0.00'


@dataclass
class Section:
    """Una tabla del reporte. `money` son los índices de columnas en USD."""

    title: str
    headers: list[str]
    rows: list[list[Any]]
    money: set[int] = field(default_factory=set)


@dataclass
class Report:
    title: str
    subtitle: str = ""
    # Pares (etiqueta, valor) que se muestran arriba, p. ej. totales.
    # Un float se muestra como dinero; un int, como conteo.
    summary: list[tuple[str, Any]] = field(default_factory=list)
    sections: list[Section] = field(default_factory=list)


def _money(value: Any) -> str:
    try:
        return f"${float(value):,.2f}"
    except (TypeError, ValueError):
        return "" if value is None else str(value)


def _text(value: Any) -> str:
    return "" if value is None else str(value)


# ------------------------------------------------------------------ Excel

def write_xlsx(report: Report, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.title[:31] or "Reporte"
    header_fill = PatternFill("solid", fgColor="123B43")
    thin = Side(style="thin", color="D5E3E1")

    sheet.append([INSTITUTION])
    sheet["A1"].font = Font(bold=True, size=10, color="13A895")
    sheet.append([report.title])
    sheet["A2"].font = Font(bold=True, size=14)
    sheet.append([report.subtitle or f"Generado el {datetime.now():%d/%m/%Y %H:%M}"])
    sheet["A3"].font = Font(italic=True, color="668087")
    sheet.append([])

    for label, value in report.summary:
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True)
        # Convención: float = dinero, int = conteo (igual que en el PDF).
        if isinstance(value, float):
            sheet.cell(sheet.max_row, 2).number_format = _MONEY_FORMAT
    if report.summary:
        sheet.append([])

    widths: dict[int, int] = {}
    for section in report.sections:
        sheet.append([section.title])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True, size=12)
        sheet.append(section.headers)
        header_row = sheet.max_row
        for col in range(1, len(section.headers) + 1):
            cell = sheet.cell(header_row, col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in section.rows:
            sheet.append([value if value is not None else "" for value in row])
            for col, value in enumerate(row, start=1):
                cell = sheet.cell(sheet.max_row, col)
                cell.border = Border(bottom=thin)
                if col - 1 in section.money and isinstance(value, (int, float)):
                    cell.number_format = _MONEY_FORMAT
        if not section.rows:
            sheet.append(["Sin registros"])
        for col, header in enumerate(section.headers, start=1):
            longest = max([len(_text(header))] + [len(_text(r[col - 1])) for r in section.rows if col - 1 < len(r)])
            widths[col] = max(widths.get(col, 10), min(longest + 2, 50))
        sheet.append([])

    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width
    workbook.save(path)


# ------------------------------------------------------------------ PDF

def report_html(report: Report) -> str:
    summary = "".join(
        f"<tr><td style='color:#668087;padding:2px 12px 2px 0'>{escape(label)}</td>"
        f"<td style='font-weight:bold'>{escape(_money(value) if isinstance(value, float) else _text(value))}</td></tr>"
        for label, value in report.summary
    )
    sections = []
    for section in report.sections:
        head = "".join(
            f"<th style='background:#123b43;color:white;padding:4px;font-size:8pt'>{escape(h)}</th>"
            for h in section.headers
        )
        body = "".join(
            "<tr>" + "".join(
                f"<td style='padding:3px;border-bottom:1px solid #d5e3e1;font-size:8pt;"
                f"{'text-align:right' if i in section.money else ''}'>"
                f"{escape(_money(v) if i in section.money else _text(v))}</td>"
                for i, v in enumerate(row)
            ) + "</tr>"
            for row in section.rows
        ) or f"<tr><td colspan='{len(section.headers)}' style='color:#668087'>Sin registros</td></tr>"
        sections.append(
            f"<h3 style='color:#123b43;margin-top:14px'>{escape(section.title)}</h3>"
            f"<table width='100%' cellspacing='0' style='border-collapse:collapse'><tr>{head}</tr>{body}</table>"
        )
    return f"""
    <html><body style="font-family:Arial;color:#183039">
      <div style="font-size:9pt;letter-spacing:2px;color:#13a895;font-weight:bold">{escape(INSTITUTION.upper())}</div>
      <h2 style="margin:4px 0">{escape(report.title)}</h2>
      <div style="color:#668087;font-size:9pt">{escape(report.subtitle or f"Generado el {datetime.now():%d/%m/%Y %H:%M}")}</div>
      <table style="margin-top:10px;font-size:9pt">{summary}</table>
      {''.join(sections)}
    </body></html>
    """


def write_pdf(report: Report, path: str) -> None:
    document = QTextDocument()
    document.setHtml(report_html(report))
    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
    printer.setOutputFileName(path)
    # Horizontal: los reportes de pagos tienen muchas columnas.
    printer.setPageLayout(QPageLayout(
        QPageSize(QPageSize.PageSizeId.Letter), QPageLayout.Orientation.Landscape,
        QMarginsF(12, 12, 12, 12), QPageLayout.Unit.Millimeter,
    ))
    document.print_(printer)


# ------------------------------------------------------------------ diálogo

def export_report(parent: QWidget, report: Report, file_stem: str) -> Optional[str]:
    """Pregunta dónde guardar (Excel o PDF según la extensión elegida) y
    escribe el archivo. Devuelve la ruta, o None si se canceló o falló."""
    folder = QStandardPaths.writableLocation(QStandardPaths.StandardLocation.DocumentsLocation)
    safe_stem = "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in file_stem).strip()
    suggested = str(Path(folder) / f"{safe_stem}_{datetime.now():%Y%m%d}.xlsx")
    path, chosen = QFileDialog.getSaveFileName(
        parent, "Exportar reporte", suggested, "Excel (*.xlsx);;PDF (*.pdf)"
    )
    if not path:
        return None
    # El formato lo decide el filtro elegido (o la extensión escrita); el
    # nombre sugerido termina en .xlsx, así que con PDF se reemplaza.
    wants_pdf = chosen.startswith("PDF") or path.lower().endswith(".pdf")
    stem = path[:-5] if path.lower().endswith(".xlsx") else path[:-4] if path.lower().endswith(".pdf") else path
    path = stem + (".pdf" if wants_pdf else ".xlsx")
    try:
        (write_pdf if wants_pdf else write_xlsx)(report, path)
    except PermissionError:
        QMessageBox.critical(parent, "No se pudo exportar", "El archivo está abierto en otro programa. Ciérralo e intenta de nuevo.")
        return None
    except Exception as exc:  # noqa: BLE001 - se informa cualquier fallo al usuario
        QMessageBox.critical(parent, "No se pudo exportar", str(exc))
        return None
    QMessageBox.information(parent, "Reporte exportado", f"Se guardó en:\n{path}")
    return path
